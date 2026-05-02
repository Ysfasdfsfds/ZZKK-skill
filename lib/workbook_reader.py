from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class RequirementEntry:
    row_number: int
    user_requirement_id: str
    rd_requirement_id: str
    sequence_id: str
    title: str
    rd_description: str
    user_description: str
    acceptance: str
    keywords: str
    priority: str
    owner_module: str
    collaborator_module: str
    customer_type: str
    usage_context: str
    scenario_name: str


@dataclass
class WorkbookData:
    workbook_path: Path
    project_name: str
    platform_branch: str
    product_name: str
    requirements: list[RequirementEntry]


@dataclass
class WorkbookPreparationResult:
    output_path: Path
    added_user_description_column: bool
    added_acceptance_column: bool
    filled_user_description_rows: int
    updated_acceptance_rows: int
    total_requirement_rows: int


@dataclass
class WorkbookHeaderInspection:
    workbook_path: Path
    missing_required_headers: list[str]
    has_user_description: bool
    has_acceptance: bool


REQUIRED_SOURCE_HEADERS = [
    '用户需求id',
    '研发需求id',
    '研发需求名称',
    '研发需求描述',
]


_OWNER_RULES: list[tuple[tuple[str, ...], str]] = [
    (("开始菜单", "搜索栏", "电源菜单", "分类切换"), "开始菜单"),
    (("工作区", "多任务"), "多任务视图"),
    (("云桌面", "dbus", "d-bus", "D-Bus", "霸屏"), "云桌面"),
    (("虚拟键盘",), "虚拟键盘"),
    (("输入法", "候选词", "快捷键"), "输入法框架"),
    (("ftp", "FTP", "文件管理", "文管", "挂载盘符", "NAS"), "文管"),
    (("主题", "神灯", "拖尾", "窗口最小化"), "设置-主题"),
    (("桌面", "小工具", "图标"), "桌面"),
    (("触控", "TrackPoint", "小红帽"), "触控设置"),
]

_COLLABORATOR_RULES: dict[str, str] = {
    '开始菜单': '电源菜单、分类面板',
    '多任务视图': '工作区管理',
    '云桌面': '会话管理、设置守护进程',
    '虚拟键盘': '主题样式、窗口渲染',
    '输入法框架': '输入法配置中心',
    '文管': 'FTP 远程连接',
    '设置-主题': '辅助功能、窗口管理',
    '桌面': '桌面组件、文件管理',
    '触控设置': '设备设置、驱动适配',
}

_CUSTOMER_TYPE_RULES: list[tuple[tuple[str, ...], str]] = [
    (("云桌面", "霸屏", "第三方客户端", "NAS"), '集成厂商/运维人员'),
    (("输入法", "文件管理", "FTP"), '桌面终端用户'),
    (("触控", "虚拟键盘", "工作区", "桌面", "主题"), '桌面终端用户'),
]

_USAGE_RULES: list[tuple[tuple[str, ...], str]] = [
    (("开始菜单",), '开始菜单高频交互场景'),
    (("工作区", "多任务"), '多任务视图与多工作区办公场景'),
    (("云桌面", "霸屏"), '云桌面单点登录与统一策略管控场景'),
    (("虚拟键盘",), '触控输入与屏幕键盘使用场景'),
    (("输入法",), '输入法个性化配置场景'),
    (("FTP", "ftp", "文件管理", "文管"), '远程 FTP 文件传输场景'),
    (("主题", "神灯", "拖尾"), '系统主题与辅助功能个性化设置场景'),
    (("触控", "TrackPoint", "小红帽"), '触控设备与外设输入场景'),
]


def inspect_workbook_headers(workbook_path: Path) -> WorkbookHeaderInspection:
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook.active
    headers = _header_map(sheet)
    missing = [header for header in REQUIRED_SOURCE_HEADERS if header not in headers]
    return WorkbookHeaderInspection(
        workbook_path=workbook_path,
        missing_required_headers=missing,
        has_user_description='用户需求描述' in headers,
        has_acceptance='验收标准' in headers,
    )


def planned_enriched_workbook_path(workbook_path: Path, output_dir: Path) -> Path:
    return output_dir / f'{workbook_path.stem}_已补充用户需求描述和验收标准.xlsx'


def enrich_workbook(workbook_path: Path, output_path: Path) -> WorkbookPreparationResult:
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    headers = _header_map(sheet)
    missing = [header for header in REQUIRED_SOURCE_HEADERS if header not in headers]
    if missing:
        missing_text = '、'.join(missing)
        raise ValueError(f'原始基础信息表必须包含这些列：{missing_text}')

    added_user_description_column = False
    added_acceptance_column = False
    if '用户需求描述' not in headers:
        insert_at = headers['验收标准'] if '验收标准' in headers else headers['研发需求描述'] + 1
        sheet.insert_cols(insert_at)
        sheet.cell(row=2, column=insert_at).value = '用户需求描述'
        added_user_description_column = True
        headers = _header_map(sheet)
    if '验收标准' not in headers:
        insert_at = headers['用户需求描述'] + 1
        sheet.insert_cols(insert_at)
        sheet.cell(row=2, column=insert_at).value = '验收标准'
        added_acceptance_column = True
        headers = _header_map(sheet)

    filled_user_description_rows = 0
    updated_acceptance_rows = 0
    total_requirement_rows = 0
    for row in range(3, sheet.max_row + 1):
        title = _cell_text(sheet.cell(row=row, column=headers['研发需求名称']).value)
        rd_description = _normalize_block(sheet.cell(row=row, column=headers['研发需求描述']).value)
        if not title and not rd_description:
            continue
        total_requirement_rows += 1

        user_desc_cell = sheet.cell(row=row, column=headers['用户需求描述'])
        acceptance_cell = sheet.cell(row=row, column=headers['验收标准'])
        current_user_description = _normalize_block(user_desc_cell.value)
        current_acceptance = _normalize_block(acceptance_cell.value)

        generated_user_description = build_user_description(title, rd_description, current_acceptance)
        generated_acceptance = build_acceptance_text(title, rd_description, current_acceptance)

        if not current_user_description and generated_user_description:
            user_desc_cell.value = generated_user_description
            filled_user_description_rows += 1
        if generated_acceptance and generated_acceptance != current_acceptance:
            acceptance_cell.value = generated_acceptance
            updated_acceptance_rows += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return WorkbookPreparationResult(
        output_path=output_path,
        added_user_description_column=added_user_description_column,
        added_acceptance_column=added_acceptance_column,
        filled_user_description_rows=filled_user_description_rows,
        updated_acceptance_rows=updated_acceptance_rows,
        total_requirement_rows=total_requirement_rows,
    )


def read_workbook(workbook_path: Path) -> WorkbookData:
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook.active
    headers = _header_map(sheet)

    required_headers = [
        '所属产品',
        '平台/分支',
        '用户需求id',
        '研发需求id',
        '编号',
        '研发需求名称',
        '研发需求描述',
        '用户需求描述',
        '验收标准',
        '关键词',
        '优先级',
    ]
    missing = [header for header in required_headers if header not in headers]
    if missing:
        raise ValueError(f"missing workbook headers: {', '.join(missing)}")

    project_name = ''
    platform_branch = ''
    product_name = ''
    requirements: list[RequirementEntry] = []

    for row in range(3, sheet.max_row + 1):
        title = _cell_text(sheet.cell(row=row, column=headers['研发需求名称']).value)
        rd_description = _normalize_block(sheet.cell(row=row, column=headers['研发需求描述']).value)
        user_description = _normalize_block(sheet.cell(row=row, column=headers['用户需求描述']).value)
        acceptance = _normalize_block(sheet.cell(row=row, column=headers['验收标准']).value)
        if not any([title, rd_description, user_description, acceptance]):
            continue

        if not project_name:
            project_name = _cell_text(sheet.cell(row=row, column=headers['所属产品']).value)
        if not platform_branch:
            platform_branch = _cell_text(sheet.cell(row=row, column=headers['平台/分支']).value)
        if not product_name:
            product_name = project_name or platform_branch

        owner_module = infer_owner_module(title, rd_description, user_description)
        collaborator_module = infer_collaborator_module(owner_module)

        requirements.append(
            RequirementEntry(
                row_number=row,
                user_requirement_id=_cell_text(sheet.cell(row=row, column=headers['用户需求id']).value),
                rd_requirement_id=_cell_text(sheet.cell(row=row, column=headers['研发需求id']).value) or _cell_text(sheet.cell(row=row, column=headers['编号']).value),
                sequence_id=_cell_text(sheet.cell(row=row, column=headers['编号']).value),
                title=title,
                rd_description=rd_description,
                user_description=user_description,
                acceptance=acceptance,
                keywords=_normalize_inline(sheet.cell(row=row, column=headers['关键词']).value),
                priority=normalize_priority(_cell_text(sheet.cell(row=row, column=headers['优先级']).value)),
                owner_module=owner_module,
                collaborator_module=collaborator_module,
                customer_type=infer_customer_type(title, rd_description, user_description),
                usage_context=infer_usage_context(title, rd_description, user_description),
                scenario_name=infer_scenario_name(title),
            )
        )

    if not requirements:
        raise ValueError(f'no usable requirement rows found in workbook: {workbook_path}')

    return WorkbookData(
        workbook_path=workbook_path,
        project_name=project_name or '银河麒麟桌面操作系统V11 SP1',
        platform_branch=platform_branch or 'UKUI4.22',
        product_name=product_name or project_name or '银河麒麟桌面操作系统V11 SP1',
        requirements=requirements,
    )


def build_user_description(title: str, rd_description: str, acceptance: str) -> str:
    owner_module = infer_owner_module(title, rd_description, '')
    value_hint = _value_hint(owner_module)
    summary = first_sentence(_strip_embedded_acceptance(rd_description)) or title
    return f'用户希望围绕“{title}”获得更稳定、清晰、可配置的{owner_module}能力，{value_hint}。具体诉求为：{summary}'


def build_acceptance_text(title: str, rd_description: str, acceptance: str) -> str:
    extracted = _extract_acceptance_lines(acceptance) or _extract_acceptance_lines(rd_description)
    if not extracted:
        extracted = [
            f'执行“{title}”对应操作时，功能可以正确触发并显示正常。',
            '关键交互、配置结果或状态反馈符合需求描述。',
            '异常场景下给出明确提示，并保持基础能力可用。',
        ]
    normalized = [_normalize_acceptance_line(line) for line in extracted if _normalize_acceptance_line(line)]
    if '异常场景下给出明确提示，并保持基础能力可用。' not in normalized:
        normalized.append('异常场景下给出明确提示，并保持基础能力可用。')
    unique_lines: list[str] = []
    seen = set()
    for line in normalized:
        if line not in seen:
            seen.add(line)
            unique_lines.append(line)
    return '\n'.join(f'{index}. {line}' for index, line in enumerate(unique_lines, start=1))


def normalize_priority(value: str) -> str:
    lowered = value.strip()
    if lowered in {'紧急', '高', 'P0', 'P1'}:
        return '高'
    if lowered in {'低', 'P3'}:
        return '低'
    return '中'


def infer_owner_module(title: str, rd_description: str, user_description: str) -> str:
    haystack = f'{title}\n{rd_description}\n{user_description}'
    for keywords, module in _OWNER_RULES:
        if any(keyword in haystack for keyword in keywords):
            return module
    return '桌面环境'


def infer_collaborator_module(owner_module: str) -> str:
    return _COLLABORATOR_RULES.get(owner_module, '无')


def infer_customer_type(title: str, rd_description: str, user_description: str) -> str:
    haystack = f'{title}\n{rd_description}\n{user_description}'
    for keywords, audience in _CUSTOMER_TYPE_RULES:
        if any(keyword in haystack for keyword in keywords):
            return audience
    return '桌面终端用户'


def infer_usage_context(title: str, rd_description: str, user_description: str) -> str:
    haystack = f'{title}\n{rd_description}\n{user_description}'
    for keywords, context in _USAGE_RULES:
        if any(keyword in haystack for keyword in keywords):
            return context
    return '桌面环境高频使用场景'


def infer_scenario_name(title: str) -> str:
    cleaned = re.sub(r'[【】]', '', title).strip()
    if not cleaned:
        return '需求场景'
    cleaned = cleaned.replace('---', '-')
    if cleaned.endswith('场景'):
        return cleaned[:24]
    return f'{cleaned[:20]}场景'


def first_sentence(text: str) -> str:
    normalized = _normalize_inline(text)
    if not normalized:
        return ''
    parts = re.split(r'(?<=[。！？!?；;])', normalized, maxsplit=1)
    return parts[0][:120]


def _header_map(sheet: Worksheet) -> dict[str, int]:
    return {
        str(sheet.cell(row=2, column=col).value).strip(): col
        for col in range(1, sheet.max_column + 1)
        if sheet.cell(row=2, column=col).value
    }


def _value_hint(owner_module: str) -> str:
    if owner_module == '云桌面':
        return '以降低集成改造成本并保证策略控制一致性'
    if owner_module in {'文管', '输入法框架'}:
        return '以减少重复操作并提升配置与使用效率'
    if owner_module in {'设置-主题', '虚拟键盘', '桌面'}:
        return '以提升界面体验一致性与个性化感受'
    return '以降低操作成本并提升日常使用体验'


def _extract_acceptance_lines(text: str) -> list[str]:
    normalized = _normalize_block(text)
    if not normalized:
        return []
    if '验收标准' in normalized:
        normalized = re.split(r'验收标准[:：】\]]*', normalized, maxsplit=1)[-1].strip()
    lines = []
    for raw_line in normalized.split('\n'):
        line = _normalize_acceptance_line(raw_line)
        if line:
            lines.append(line)
    return lines


def _normalize_acceptance_line(text: str) -> str:
    line = re.sub(r'^[\s\-•·]+', '', text).strip()
    line = re.sub(r'^[0-9]+[\.|、\)]\s*', '', line).strip()
    line = re.sub(r'^（?[0-9]+）\s*', '', line).strip()
    return line


def _strip_embedded_acceptance(text: str) -> str:
    normalized = _normalize_block(text)
    if '验收标准' in normalized:
        normalized = re.split(r'验收标准[:：】\]]*', normalized, maxsplit=1)[0].strip()
    return normalized


def _cell_text(value: object) -> str:
    if value is None:
        return ''
    return str(value).strip()


def _normalize_inline(value: object) -> str:
    text = _cell_text(value)
    return re.sub(r'\s+', ' ', text).strip()


def _normalize_block(value: object) -> str:
    text = _cell_text(value).replace('\r', '\n')
    lines = [re.sub(r'\s+', ' ', line).strip() for line in text.split('\n')]
    return '\n'.join(line for line in lines if line)
