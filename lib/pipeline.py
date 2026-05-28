from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W_NS}

from lib.doc_numbering import (
    customer_doc_filename,
    customer_doc_number,
    customer_review_filename,
    customer_review_record_number,
    dotted_date,
    ensure_run_date,
    product_doc_filename,
    product_doc_number,
    product_review_filename,
    product_review_record_number,
    version_token,
)
from lib.docx_ops import (
    CustomerDocData,
    CustomerRequirementRow,
    CustomerScenario,
    ProductDocData,
    ProductRequirement,
    render_customer_doc,
    render_product_doc,
)
from lib.review_logic import (
    build_customer_review_plan,
    build_product_review_plan,
    customer_requirement_description as review_customer_requirement_description,
    customer_scenario_description as review_customer_scenario_description,
    product_quality_rows,
    product_requirement_details,
)
from lib.workbook_reader import (
    WorkbookData,
    enrich_workbook,
    inspect_workbook_headers,
    planned_enriched_workbook_path,
    read_workbook,
)
from lib.xlsx_review_sheet_ops import render_customer_review_sheet, render_product_review_sheet


@dataclass
class PipelineConfig:
    workbook: Path
    customer_template: Path
    product_template: Path
    review_template: Path
    module_code: str
    module_name: str
    product_line_code: str
    product_name: str
    draft_version: str
    final_version: str
    customer_draft_date: str
    customer_review_date: str
    customer_final_date: str
    product_draft_date: str
    product_review_date: str
    product_final_date: str
    output_dir: Path
    module_product_manager: str
    module_project_manager: str
    reviewers: str
    rd_owner: str
    product_owner: str
    test_owner: str
    security_owner: str
    qa: str
    run_date: str | None = None
    overwrite: bool = False


def plan_outputs(config: PipelineConfig) -> dict[str, object]:
    run_date = ensure_run_date(config.run_date)
    customer_review_date_token = ensure_run_date(config.customer_review_date)
    product_review_date_token = ensure_run_date(config.product_review_date)
    _validate_inputs(config)
    outputs = _output_paths(config)
    workbook_check = inspect_workbook_headers(config.workbook)
    prepared_workbook = planned_enriched_workbook_path(config.workbook, config.output_dir)
    if workbook_check.missing_required_headers:
        missing_text = '、'.join(workbook_check.missing_required_headers)
        raise ValueError(f'原始基础信息表必须包含这些列：{missing_text}')
    return {
        'module_code': config.module_code,
        'module_name': config.module_name,
        'product_line_code': config.product_line_code,
        'product_name': config.product_name,
        'draft_version': config.draft_version,
        'final_version': config.final_version,
        'run_date': run_date,
        'customer_document_dates': {
            'customer_draft_date': config.customer_draft_date,
            'customer_review_date': config.customer_review_date,
            'customer_final_date': config.customer_final_date,
            'date_format_rule': '按用户输入原样写入，不做格式转换',
        },
        'product_document_dates': {
            'product_draft_date': config.product_draft_date,
            'product_review_date': config.product_review_date,
            'product_final_date': config.product_final_date,
            'date_format_rule': '按用户输入原样写入，不做格式转换',
        },
        'required_source_headers': ['用户需求id', '研发需求id', '研发需求名称', '研发需求描述'],
        'review_record_people': {
            'host': config.module_product_manager,
            'scribe': config.module_project_manager,
            'reviewers': config.reviewers,
            'rd_owner': config.rd_owner,
            'product_owner': config.product_owner,
            'test_owner': config.test_owner,
            'security_owner': config.security_owner,
            'other_people': config.qa,
        },
        'doc_number_rules_from_templates': {
            'customer_doc': '产线名称英文缩写-模块ID-CRS-版本号',
            'product_doc': '产线名称英文缩写-产品名称-模块ID-PRD-版本号',
        },
        'workbook_preparation': {
            'source_workbook': str(config.workbook),
            'prepared_workbook': str(prepared_workbook),
            'has_user_description_column': workbook_check.has_user_description,
            'has_acceptance_column': workbook_check.has_acceptance,
            'will_add_or_fill': ['用户需求描述', '验收标准'],
        },
        'doc_numbers': {
            'customer_draft': customer_doc_number(config.module_code, config.draft_version, config.product_line_code, config.product_name),
            'customer_final': customer_doc_number(config.module_code, config.final_version, config.product_line_code, config.product_name),
            'product_draft': product_doc_number(config.module_code, config.draft_version, config.product_line_code, config.product_name),
            'product_final': product_doc_number(config.module_code, config.final_version, config.product_line_code, config.product_name),
            'customer_review': customer_review_record_number(config.module_code, customer_review_date_token),
            'product_review': product_review_record_number(config.module_code, product_review_date_token),
        },
        'content_correspondence': {
            'generation_flow': [
                'customer_draft',
                'customer_review',
                'customer_final',
                'product_draft',
                'product_review',
                'product_final',
            ],
            'customer_review_targets': outputs['customer_draft'].stem,
            'customer_final_must_match_requirement_scope_of': outputs['customer_draft'].stem,
            'product_review_targets': outputs['product_draft'].stem,
            'product_final_must_match_requirement_scope_of': outputs['product_draft'].stem,
        },
        'semantic_review': {
            'review_sheet_must_not_default_to_no_change': True,
            'customer_review_perspectives': ['模块产品经理', '产品级产品经理', '安全负责人', '测试负责人'],
            'product_review_perspectives': ['项目经理', '研发负责人', '安全负责人', '测试负责人'],
            'final_docs_are_revised_from_review_issues': True,
        },
        'outputs': {name: str(path) for name, path in outputs.items()},
        'existing_outputs': {name: path.exists() for name, path in outputs.items()},
    }


def run_pipeline(config: PipelineConfig) -> dict[str, object]:
    run_date = ensure_run_date(config.run_date)
    customer_review_date_token = ensure_run_date(config.customer_review_date)
    product_review_date_token = ensure_run_date(config.product_review_date)
    _validate_inputs(config)
    outputs = _output_paths(config)
    prepared_workbook_path = planned_enriched_workbook_path(config.workbook, config.output_dir)
    collisions = [str(path) for path in [prepared_workbook_path, *outputs.values()] if path.exists()]
    if collisions and not config.overwrite:
        raise FileExistsError(f'output files already exist: {collisions}')

    config.output_dir.mkdir(parents=True, exist_ok=True)
    preparation = enrich_workbook(config.workbook, prepared_workbook_path)
    workbook = read_workbook(preparation.output_path)

    customer_draft = _build_customer_doc_data(config, workbook, config.draft_version, run_date, final=False)
    product_draft = _build_product_doc_data(config, workbook, config.draft_version, run_date, final=False)

    render_customer_doc(config.customer_template, outputs['customer_draft'], customer_draft)
    render_product_doc(config.product_template, outputs['product_draft'], product_draft)

    customer_review_plan = build_customer_review_plan(
        module_name=config.module_name,
        draft_title=outputs['customer_draft'].stem,
        final_version=config.final_version,
        workbook=workbook,
        module_product_manager=config.module_product_manager,
        product_owner=config.product_owner,
        test_owner=config.test_owner,
        security_owner=config.security_owner,
    )
    product_review_plan = build_product_review_plan(
        module_name=config.module_name,
        draft_title=outputs['product_draft'].stem,
        final_version=config.final_version,
        workbook=workbook,
        module_project_manager=config.module_project_manager,
        rd_owner=config.rd_owner,
        test_owner=config.test_owner,
        security_owner=config.security_owner,
    )

    render_customer_review_sheet(
        config.review_template,
        outputs['customer_review'],
        record_number=customer_review_record_number(config.module_code, customer_review_date_token),
        project_name=_review_product_name(config),
        work_product_name=_review_product_name(config),
        work_product_review_scope=f'《{outputs["customer_draft"].stem}》',
        meeting_date=config.customer_review_date,
        host=config.module_product_manager,
        scribe=config.module_project_manager,
        reviewers=config.reviewers,
        other_people=config.qa,
        review_plan=customer_review_plan,
        revision_reference=f'《{outputs["customer_draft"].stem}》',
    )

    customer_final = _build_customer_doc_data(config, workbook, config.final_version, customer_review_date_token, final=True)
    product_final = _build_product_doc_data(config, workbook, config.final_version, product_review_date_token, final=True)
    render_customer_doc(config.customer_template, outputs['customer_final'], customer_final)
    render_product_doc(config.product_template, outputs['product_final'], product_final)
    render_product_review_sheet(
        config.review_template,
        outputs['product_review'],
        record_number=product_review_record_number(config.module_code, product_review_date_token),
        project_name=_review_product_name(config),
        work_product_name=_review_product_name(config),
        work_product_review_scope=f'《{outputs["product_draft"].stem}》',
        meeting_date=config.product_review_date,
        host=config.module_product_manager,
        scribe=config.module_project_manager,
        reviewers=config.reviewers,
        other_people=config.qa,
        review_plan=product_review_plan,
        revision_reference=f'《{outputs["product_draft"].stem}》',
    )

    verification = verify_outputs(config)
    return {
        'run_date': run_date,
        'prepared_workbook': {
            'path': str(preparation.output_path),
            'added_user_description_column': preparation.added_user_description_column,
            'added_acceptance_column': preparation.added_acceptance_column,
            'filled_user_description_rows': preparation.filled_user_description_rows,
            'updated_acceptance_rows': preparation.updated_acceptance_rows,
            'total_requirement_rows': preparation.total_requirement_rows,
        },
        'generated_files': {name: str(path) for name, path in outputs.items()},
        'verification': verification,
    }


def verify_outputs(config: PipelineConfig) -> dict[str, object]:
    run_date = ensure_run_date(config.run_date)
    customer_review_date_token = ensure_run_date(config.customer_review_date)
    product_review_date_token = ensure_run_date(config.product_review_date)
    outputs = _output_paths(config)
    strict_review_date = config.run_date is not None
    expected_strings = {
        'customer_draft': customer_doc_number(config.module_code, config.draft_version, config.product_line_code, config.product_name),
        'customer_final': customer_doc_number(config.module_code, config.final_version, config.product_line_code, config.product_name),
        'product_draft': product_doc_number(config.module_code, config.draft_version, config.product_line_code, config.product_name),
        'product_final': product_doc_number(config.module_code, config.final_version, config.product_line_code, config.product_name),
        'customer_review': customer_review_record_number(config.module_code, customer_review_date_token),
        'product_review': product_review_record_number(config.module_code, product_review_date_token),
    }
    review_prefixes = {
        'customer_review': f'MT-PR-A-{config.module_code}-CRS-',
        'product_review': f'MT-PR-A-{config.module_code}-PRD-',
    }
    results: dict[str, dict[str, object]] = {}
    for name, path in outputs.items():
        item = {
            'path': str(path),
            'exists': path.exists(),
            'expected_identifier': expected_strings[name],
            'identifier_present': False,
        }
        if path.exists() and path.suffix == '.docx':
            text = _read_docx_text(path)
            item['identifier_present'] = expected_strings[name] in text
        elif path.exists() and path.suffix == '.xlsx':
            text = _read_xlsx_text(path)
            if strict_review_date:
                item['identifier_present'] = expected_strings[name] in text
            else:
                item['identifier_present'] = review_prefixes[name] in text
        results[name] = item
    correspondence = _verify_content_correspondence(config, outputs)
    return {
        'run_date': run_date,
        'all_outputs_present': all(item['exists'] for item in results.values()),
        'all_identifiers_present': all(item['identifier_present'] for item in results.values()),
        'all_content_correspondence_passed': correspondence['passed'],
        'content_correspondence': correspondence,
        'outputs': results,
    }


def _validate_inputs(config: PipelineConfig) -> None:
    for path in [config.workbook, config.customer_template, config.product_template, config.review_template]:
        if not path.exists():
            raise FileNotFoundError(path)
    required_text_fields = {
        'module_code': config.module_code,
        'module_name': config.module_name,
        'product_line_code': config.product_line_code,
        'product_name': config.product_name,
        'draft_version': config.draft_version,
        'final_version': config.final_version,
        'customer_draft_date': config.customer_draft_date,
        'customer_review_date': config.customer_review_date,
        'customer_final_date': config.customer_final_date,
        'product_draft_date': config.product_draft_date,
        'product_review_date': config.product_review_date,
        'product_final_date': config.product_final_date,
        'module_product_manager': config.module_product_manager,
        'module_project_manager': config.module_project_manager,
        'reviewers': config.reviewers,
        'rd_owner': config.rd_owner,
        'product_owner': config.product_owner,
        'test_owner': config.test_owner,
        'security_owner': config.security_owner,
        'qa': config.qa,
    }
    missing = [name for name, value in required_text_fields.items() if not value.strip()]
    if missing:
        raise ValueError(f"missing required input fields: {', '.join(missing)}")


def _output_paths(config: PipelineConfig) -> dict[str, Path]:
    return {
        'customer_draft': config.output_dir / customer_doc_filename(config.module_name, config.draft_version),
        'customer_final': config.output_dir / customer_doc_filename(config.module_name, config.final_version),
        'customer_review': config.output_dir / customer_review_filename(config.module_name),
        'product_draft': config.output_dir / product_doc_filename(config.module_name, config.draft_version),
        'product_final': config.output_dir / product_doc_filename(config.module_name, config.final_version),
        'product_review': config.output_dir / product_review_filename(config.module_name),
    }


def _build_customer_doc_data(config: PipelineConfig, workbook: WorkbookData, version: str, run_date: str, *, final: bool) -> CustomerDocData:
    version_display = version_token(version)
    scenarios = []
    requirement_rows = []
    for index, requirement in enumerate(workbook.requirements, start=1):
        scenario_description = review_customer_scenario_description(requirement, final=final)
        if final and index == 1:
            scenario_description = (
                f'{scenario_description} '
                f'评审追踪：根据同级评审记录 {customer_review_record_number(config.module_code, run_date)} 由初稿形成终稿。'
            )
        scenarios.append(
            CustomerScenario(
                scenario_name=requirement.scenario_name,
                customer_type=requirement.customer_type,
                usage_context=requirement.usage_context,
                scenario_id=f'S-{index:02d}',
                description=scenario_description,
            )
        )
        requirement_rows.append(
            CustomerRequirementRow(
                user_requirement_id=requirement.user_requirement_id,
                title=requirement.title,
                description=review_customer_requirement_description(requirement, scenario_id=f'S-{index:02d}', final=final),
                priority=requirement.priority,
                owner_module=config.module_name if final else requirement.owner_module,
                collaborator_module=_customer_collaborator(requirement, final=final),
            )
        )

    return CustomerDocData(
        module_name=config.module_name,
        doc_number=customer_doc_number(config.module_code, version, config.product_line_code, config.product_name),
        version_display=version_display,
        approval_rows=_approval_rows(
            config,
            final=final,
            draft_date=config.customer_draft_date,
            review_date=config.customer_review_date,
            final_date=config.customer_final_date,
        ),
        version_rows=_version_rows(
            config,
            final=final,
            draft_date=config.customer_draft_date,
            final_date=config.customer_final_date,
        ),
        scenarios=scenarios,
        requirement_rows=requirement_rows,
    )


def _build_product_doc_data(config: PipelineConfig, workbook: WorkbookData, version: str, run_date: str, *, final: bool) -> ProductDocData:
    version_display = version_token(version)
    requirements = []
    for requirement in workbook.requirements:
        details = product_requirement_details(requirement, final=final)
        requirements.append(
            ProductRequirement(
                module_name=config.module_name,
                level='C1',
                related_module_name=details['related_module_name'],
                parent_module_name='无',
                chipset_special='无',
                user_requirement_id=requirement.user_requirement_id,
                rd_requirement_id=requirement.rd_requirement_id or requirement.sequence_id,
                priority_marks=_priority_marks(requirement.priority),
                implement_way='☑ 自研   □ 开源修改 □ 开源引入或第三方',
                feature_description=details['feature_description'],
                related_description=details['related_description'],
                acceptance=details['acceptance'],
                input_desc=details['input_desc'],
                process_desc=details['process_desc'],
                output_desc=details['output_desc'],
                exception_desc=details['exception_desc'],
                layout_desc=details['layout_desc'],
                quality_desc=details['quality_desc'],
            )
        )

    quality_rows = product_quality_rows(config.module_name, final=final, requirements=workbook.requirements)
    reference_rows = _product_reference_rows(config, run_date, final=final)

    return ProductDocData(
        module_name=config.module_name,
        doc_number=product_doc_number(config.module_code, version, config.product_line_code, config.product_name),
        version_display=version_display,
        run_date_display=dotted_date(run_date),
        approval_rows=_approval_rows(
            config,
            final=final,
            draft_date=config.product_draft_date,
            review_date=config.product_review_date,
            final_date=config.product_final_date,
        ),
        version_rows=_version_rows(
            config,
            final=final,
            draft_date=config.product_draft_date,
            final_date=config.product_final_date,
        ),
        module_description=f'{config.module_name}模块属于桌面环境核心组成部分，负责承接本轮需求中与终端界面、交互入口、配置体验及集成能力相关的产品落地。',
        tech_constraints='需遵循现有桌面环境、设置框架和接口约束，优先复用既有实现方式。',
        resource_constraints='需结合当前版本节奏与研发资源推进，优先保证高价值需求按期落地。',
        assumptions='默认目标环境具备桌面环境基础能力，上下游模块可提供必要接口、配置与联调支持。',
        requirements=requirements,
        quality_rows=quality_rows,
        glossary_rows=_glossary_rows(workbook),
        reference_rows=reference_rows,
    )


def _customer_collaborator(requirement, *, final: bool) -> str:
    if not final:
        return requirement.collaborator_module
    return product_requirement_details(requirement, final=True)['related_module_name']


def _approval_rows(
    config: PipelineConfig,
    *,
    final: bool,
    draft_date: str,
    review_date: str,
    final_date: str,
) -> list[tuple[str, str]]:
    if final:
        return [
            (config.module_product_manager, draft_date),
            (config.reviewers, review_date),
            (config.module_project_manager, final_date),
        ]
    return [
        (config.module_product_manager, draft_date),
        ('', ''),
        ('', ''),
    ]


def _version_rows(
    config: PipelineConfig,
    *,
    final: bool,
    draft_date: str,
    final_date: str,
) -> list[tuple[str, str, str, str]]:
    draft_row = (
        draft_date,
        version_token(config.draft_version),
        '初稿',
        config.module_product_manager,
    )
    if not final:
        return [draft_row]
    return [
        draft_row,
        (
            final_date,
            version_token(config.final_version),
            '终稿',
            config.module_product_manager,
        ),
    ]


def _review_product_name(config: PipelineConfig) -> str:
    name = config.product_name.strip()
    if name.startswith('银河麒麟'):
        return name
    return f'银河麒麟桌面操作系统{name}'


def _product_reference_rows(config: PipelineConfig, run_date: str, *, final: bool) -> list[tuple[str, str]]:
    prepared_workbook = planned_enriched_workbook_path(config.workbook, config.output_dir).name
    if final:
        return [
            ('输入文档', f'《{config.module_name}-客户需求说明书_{config.final_version}》'),
            ('依赖的接口契约', '按基础信息表中的接口清单、Demo、设计稿、策略说明和研发补充资料持续维护。'),
            ('遵循的标准与规范', '项目既有研发与评审规范'),
            ('同级评审记录', f'《{config.module_name}-模块级产品需求分析说明书_同级评审会议记录表》（记录编号：{product_review_record_number(config.module_code, run_date)}）'),
            ('输入基础表', prepared_workbook),
        ]
    return [
        ('输入文档', f'《{config.module_name}-客户需求说明书_{config.draft_version}》'),
        ('依赖的接口契约', '待研发补充接口清单、Demo、设计稿或策略说明。'),
        ('遵循的标准与规范', '项目既有研发与评审规范'),
        ('同级评审状态', '初稿待同级评审'),
        ('输入基础表', prepared_workbook),
    ]


def _priority_marks(priority: str) -> str:
    if priority == '高':
        return '☑ 高 □ 中 □ 低'
    if priority == '低':
        return '□ 高 □ 中 ☑ 低'
    return '□ 高 ☑ 中 □ 低'


def _glossary_rows(workbook: WorkbookData) -> list[tuple[str, str]]:
    haystack = '\n'.join(
        '\n'.join(
            [
                requirement.title,
                requirement.rd_description,
                requirement.user_description,
                requirement.acceptance,
                requirement.keywords,
            ]
        )
        for requirement in workbook.requirements
    )
    rows: list[tuple[str, str]] = []
    glossary_rules = [
        (('ki18n', 'KI18N'), 'ki18n', '银河麒麟桌面操作系统中的多语言基础组件，本次需求要求对其开展重构优化，并保证重构后通用能力无缺失、无偏差。'),
        (('国际化', 'i18n', 'I18N'), '国际化（i18n）', '软件在架构、资源和接口层面支持多语言、多地区适配的能力，是本次多语言组件重构的核心目标之一。'),
        (('本地化', 'l10n', 'L10N'), '本地化（l10n）', '针对具体语言、地区和使用习惯进行翻译、格式、排序、区域数据等适配的过程。'),
        (('本地化文本', '文本处理', '多语言文本'), '本地化文本处理', '对多语言字符串、翻译资源、显示文本和相关格式进行加载、转换、展示与一致性处理的能力。'),
        (('地域', '地区', '时区'), '地域与时区数据', '用于支持不同地区日期、时间、区域格式、时区信息等展示与处理的数据集合。'),
        (('基础组件', '操作系统基础组件'), '操作系统基础组件', '为上层桌面应用和系统能力提供公共支撑的底层组件，本次需求关注其安全性、稳定性和可替换能力。'),
        (('供应链安全',), '供应链安全', '通过组件重构、来源可控和能力替代降低外部依赖风险，提升基础组件交付和维护的安全保障能力。'),
        (('自主可控',), '自主可控', '关键基础组件具备可维护、可演进和可替代能力，减少对不可控外部实现的依赖。'),
        (('性能',), '性能', '组件在多语言资源加载、地域数据处理和文本处理等场景下的响应效率与资源占用表现。'),
        (('稳定性',), '稳定性', '组件在持续运行、异常输入和多场景调用下保持功能结果一致、不中断核心能力的质量属性。'),
    ]
    for keywords, term, definition in glossary_rules:
        if any(keyword in haystack for keyword in keywords):
            rows.append((term, definition))

    for requirement in workbook.requirements:
        if len(rows) >= 10:
            break
        if requirement.title:
            rows.append((requirement.title, '本次基础信息表中的研发需求名称，用于界定当前产品需求分析说明书的需求范围。'))

    deduped: list[tuple[str, str]] = []
    used = set()
    for key, value in rows:
        if key and key not in used:
            used.add(key)
            deduped.append((key, value))
    return deduped[:10]


def _read_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
    return ''.join(node.text or '' for node in root.findall('.//w:t', {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}))


def _read_xlsx_text(path: Path) -> str:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    values: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                values.append(str(cell.value))
    return '\n'.join(values)


def _verify_content_correspondence(config: PipelineConfig, outputs: dict[str, Path]) -> dict[str, object]:
    prepared_workbook = planned_enriched_workbook_path(config.workbook, config.output_dir)
    result: dict[str, object] = {
        'passed': False,
        'prepared_workbook': str(prepared_workbook),
        'prepared_workbook_exists': prepared_workbook.exists(),
        'customer_review_targets_draft': False,
        'customer_review_targets_final': False,
        'product_review_targets_draft': False,
        'product_review_targets_final': False,
        'customer_final_mentions_review_record': False,
        'product_final_mentions_review_record': False,
        'customer_draft_date_present': False,
        'customer_final_dates_present': False,
        'customer_review_fields_present': False,
        'product_draft_date_present': False,
        'product_final_dates_present': False,
        'product_review_fields_present': False,
        'review_sheets_avoid_placeholder_conclusions': False,
        'layout_checks_passed': False,
        'layout_checks': {},
        'missing_customer_requirement_ids_in_draft': [],
        'missing_customer_requirement_ids_in_final': [],
        'missing_product_requirement_ids_in_draft': [],
        'missing_product_requirement_ids_in_final': [],
    }
    required_outputs = [
        outputs['customer_draft'],
        outputs['customer_final'],
        outputs['customer_review'],
        outputs['product_draft'],
        outputs['product_final'],
        outputs['product_review'],
    ]
    if not prepared_workbook.exists() or not all(path.exists() for path in required_outputs):
        return result

    try:
        workbook = read_workbook(prepared_workbook)
    except Exception as exc:  # pragma: no cover - returned for CLI diagnostics
        result['error'] = str(exc)
        return result

    customer_ids = [
        requirement.user_requirement_id
        for requirement in workbook.requirements
        if requirement.user_requirement_id
    ]
    product_ids = [
        requirement.rd_requirement_id or requirement.sequence_id
        for requirement in workbook.requirements
        if requirement.rd_requirement_id or requirement.sequence_id
    ]
    result['requirement_count'] = len(workbook.requirements)

    customer_draft_text = _read_docx_text(outputs['customer_draft'])
    customer_final_text = _read_docx_text(outputs['customer_final'])
    product_draft_text = _read_docx_text(outputs['product_draft'])
    product_final_text = _read_docx_text(outputs['product_final'])
    customer_review_text = _read_xlsx_text(outputs['customer_review'])
    product_review_text = _read_xlsx_text(outputs['product_review'])

    result['customer_review_targets_draft'] = outputs['customer_draft'].stem in customer_review_text
    result['customer_review_targets_final'] = outputs['customer_final'].stem in customer_review_text
    result['product_review_targets_draft'] = outputs['product_draft'].stem in product_review_text
    result['product_review_targets_final'] = outputs['product_final'].stem in product_review_text
    result['customer_final_mentions_review_record'] = f'MT-PR-A-{config.module_code}-CRS-' in customer_final_text
    result['product_final_mentions_review_record'] = f'MT-PR-A-{config.module_code}-PRD-' in product_final_text
    result['customer_draft_date_present'] = config.customer_draft_date in customer_draft_text
    result['customer_final_dates_present'] = all(
        value in customer_final_text
        for value in [config.customer_draft_date, config.customer_review_date, config.customer_final_date]
    )
    result['customer_review_fields_present'] = all(
        value in customer_review_text
        for value in [outputs['customer_draft'].stem, config.customer_review_date]
    )
    result['product_draft_date_present'] = config.product_draft_date in product_draft_text
    result['product_final_dates_present'] = all(
        value in product_final_text
        for value in [config.product_draft_date, config.product_review_date, config.product_final_date]
    )
    result['product_review_fields_present'] = all(
        value in product_review_text
        for value in [_review_product_name(config), f'《{outputs["product_draft"].stem}》', '设计阶段', config.product_review_date]
    )
    forbidden_review_phrases = ['无需整改', '无需处理', '无待解决问题', '无遗留待解决问题']
    result['review_sheets_avoid_placeholder_conclusions'] = not any(
        phrase in customer_review_text or phrase in product_review_text
        for phrase in forbidden_review_phrases
    )

    result['missing_customer_requirement_ids_in_draft'] = _missing_tokens(customer_draft_text, customer_ids)
    result['missing_customer_requirement_ids_in_final'] = _missing_tokens(customer_final_text, customer_ids)
    result['missing_product_requirement_ids_in_draft'] = _missing_tokens(product_draft_text, product_ids)
    result['missing_product_requirement_ids_in_final'] = _missing_tokens(product_final_text, product_ids)
    result['layout_checks'] = _verify_docx_layout_structure(config, outputs, len(workbook.requirements))
    result['layout_checks_passed'] = all(result['layout_checks'].values())

    result['passed'] = (
        result['customer_review_targets_draft']
        and not result['customer_review_targets_final']
        and result['product_review_targets_draft']
        and not result['product_review_targets_final']
        and result['customer_final_mentions_review_record']
        and result['product_final_mentions_review_record']
        and result['customer_draft_date_present']
        and result['customer_final_dates_present']
        and result['customer_review_fields_present']
        and result['product_draft_date_present']
        and result['product_final_dates_present']
        and result['product_review_fields_present']
        and result['review_sheets_avoid_placeholder_conclusions']
        and result['layout_checks_passed']
        and not result['missing_customer_requirement_ids_in_draft']
        and not result['missing_customer_requirement_ids_in_final']
        and not result['missing_product_requirement_ids_in_draft']
        and not result['missing_product_requirement_ids_in_final']
    )
    return result


def _missing_tokens(text: str, tokens: list[str]) -> list[str]:
    return [token for token in tokens if token and token not in text]


def _verify_docx_layout_structure(config: PipelineConfig, outputs: dict[str, Path], requirement_count: int) -> dict[str, bool]:
    return {
        'customer_draft_version_rows': _docx_version_row_count(outputs['customer_draft']) == 1,
        'customer_final_version_rows': _docx_version_row_count(outputs['customer_final']) == 2,
        'product_draft_version_rows': _docx_version_row_count(outputs['product_draft']) == 1,
        'product_final_version_rows': _docx_version_row_count(outputs['product_final']) == 2,
        'customer_draft_version_table_preserves_template_rows': _docx_version_data_row_count(outputs['customer_draft']) >= 1,
        'customer_final_version_table_preserves_template_rows': _docx_version_data_row_count(outputs['customer_final']) >= 2,
        'product_draft_version_table_preserves_template_rows': _docx_version_data_row_count(outputs['product_draft']) >= 1,
        'product_final_version_table_preserves_template_rows': _docx_version_data_row_count(outputs['product_final']) >= 2,
        'customer_draft_version_rows_renderable': _docx_version_rows_without_vertical_merge(outputs['customer_draft']),
        'customer_final_version_rows_renderable': _docx_version_rows_without_vertical_merge(outputs['customer_final']),
        'product_draft_version_rows_renderable': _docx_version_rows_without_vertical_merge(outputs['product_draft']),
        'product_final_version_rows_renderable': _docx_version_rows_without_vertical_merge(outputs['product_final']),
        'product_draft_no_extra_requirement_tables': _docx_direct_table_count(outputs['product_draft'], ('模块名称', '功能描述：', '验收标准：')) == requirement_count,
        'product_final_no_extra_requirement_tables': _docx_direct_table_count(outputs['product_final'], ('模块名称', '功能描述：', '验收标准：')) == requirement_count,
        'product_draft_module_description_not_pushed_by_blanks': _docx_blank_paragraphs_before_heading(outputs['product_draft'], '模块描述') <= 1,
        'product_final_module_description_not_pushed_by_blanks': _docx_blank_paragraphs_before_heading(outputs['product_final'], '模块描述') <= 1,
        'product_draft_nonfunctional_not_pushed_by_blanks': _docx_blank_paragraphs_before_heading(outputs['product_draft'], '模块新增非功能详述') <= 1,
        'product_final_nonfunctional_not_pushed_by_blanks': _docx_blank_paragraphs_before_heading(outputs['product_final'], '模块新增非功能详述') <= 1,
    }


def _docx_root(path: Path) -> ET.Element:
    with zipfile.ZipFile(path) as archive:
        return ET.fromstring(archive.read('word/document.xml'))


def _docx_direct_body_children(path: Path) -> list[ET.Element]:
    body = _docx_root(path).find('w:body', NS)
    return list(body) if body is not None else []


def _docx_element_text(element: ET.Element) -> str:
    return ''.join(node.text or '' for node in element.findall('.//w:t', NS)).strip()


def _docx_version_row_count(path: Path) -> int:
    root = _docx_root(path)
    for table in root.findall('.//w:tbl', NS):
        text = _docx_element_text(table)
        if not all(marker in text for marker in ['日期', '版本号', '发布说明', '编写者']):
            continue
        count = 0
        for row in table.findall('w:tr', NS)[1:]:
            row_text = _docx_element_text(row)
            if '初稿' in row_text or '终稿' in row_text:
                count += 1
        return count
    return 0


def _docx_version_data_row_count(path: Path) -> int:
    root = _docx_root(path)
    for table in root.findall('.//w:tbl', NS):
        text = _docx_element_text(table)
        if all(marker in text for marker in ['日期', '版本号', '发布说明', '编写者']):
            return max(0, len(table.findall('w:tr', NS)) - 1)
    return 0


def _docx_version_rows_without_vertical_merge(path: Path) -> bool:
    root = _docx_root(path)
    for table in root.findall('.//w:tbl', NS):
        text = _docx_element_text(table)
        if not all(marker in text for marker in ['日期', '版本号', '发布说明', '编写者']):
            continue
        for row in table.findall('w:tr', NS):
            if not _docx_element_text(row):
                continue
            for cell in row.findall('w:tc', NS):
                props = cell.find('w:tcPr', NS)
                if props is not None and props.find('w:vMerge', NS) is not None:
                    return False
        return True
    return False


def _docx_direct_table_count(path: Path, markers: tuple[str, ...]) -> int:
    return sum(
        1
        for child in _docx_direct_body_children(path)
        if child.tag == f'{{{W_NS}}}tbl' and all(marker in _docx_element_text(child) for marker in markers)
    )


def _docx_blank_paragraphs_before_heading(path: Path, heading: str) -> int:
    children = _docx_direct_body_children(path)
    heading_index = next(
        (
            index
            for index, child in enumerate(children)
            if child.tag == f'{{{W_NS}}}p' and _docx_element_text(child) == heading
        ),
        None,
    )
    if heading_index is None:
        return 0
    count = 0
    for child in reversed(children[:heading_index]):
        if child.tag == f'{{{W_NS}}}p' and _docx_element_text(child) == '':
            count += 1
            continue
        break
    return count
