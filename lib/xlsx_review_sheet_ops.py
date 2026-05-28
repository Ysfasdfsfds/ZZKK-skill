from __future__ import annotations

from pathlib import Path

from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from lib.review_logic import ReviewPlan


_WRAP_TOP = Alignment(wrap_text=True, vertical='top')


def render_customer_review_sheet(
    template_path: Path,
    output_path: Path,
    *,
    record_number: str,
    project_name: str,
    work_product_title: str,
    meeting_date: str,
    host: str,
    scribe: str,
    reviewers: str,
    other_people: str,
    review_plan: ReviewPlan,
) -> None:
    workbook = load_workbook(template_path)
    sheet = workbook.active
    sheet.cell(row=2, column=6).value = record_number
    sheet.cell(row=3, column=5).value = project_name
    sheet.cell(row=4, column=5).value = f'《{work_product_title}》'
    sheet.cell(row=5, column=5).value = '应用场景表、用户需求描述表、需求条目、优先级、主责/协作模块及文档编号一致性检查'
    sheet.cell(row=5, column=11).value = '需求分析与文档评审'
    sheet.cell(row=6, column=5).value = meeting_date
    sheet.cell(row=6, column=9).value = 1.5
    sheet.cell(row=7, column=5).value = host
    sheet.cell(row=7, column=9).value = scribe
    sheet.cell(row=8, column=5).value = reviewers
    sheet.cell(row=9, column=5).value = other_people
    _write_review_plan(sheet, review_plan, owner=host, completion_date=meeting_date)
    workbook.save(output_path)


def render_product_review_sheet(
    template_path: Path,
    output_path: Path,
    *,
    record_number: str,
    project_name: str,
    work_product_name: str,
    work_product_review_scope: str,
    meeting_date: str,
    host: str,
    scribe: str,
    reviewers: str,
    other_people: str,
    review_plan: ReviewPlan,
) -> None:
    workbook = load_workbook(template_path)
    sheet = workbook.active
    sheet.cell(row=2, column=6).value = record_number
    sheet.cell(row=3, column=5).value = project_name
    sheet.cell(row=4, column=5).value = work_product_name
    sheet.cell(row=5, column=5).value = work_product_review_scope
    sheet.cell(row=5, column=11).value = '设计阶段'
    sheet.cell(row=6, column=5).value = meeting_date
    sheet.cell(row=6, column=9).value = 2
    sheet.cell(row=7, column=5).value = host
    sheet.cell(row=7, column=9).value = scribe
    sheet.cell(row=8, column=5).value = reviewers
    sheet.cell(row=9, column=5).value = other_people
    _write_review_plan(sheet, review_plan, owner=host, completion_date=meeting_date)
    workbook.save(output_path)


def _write_review_plan(sheet, review_plan: ReviewPlan, *, owner: str, completion_date: str) -> None:
    _clear_issue_rows(sheet)
    for offset, issue in enumerate(review_plan.issues):
        row = _issue_row(sheet, offset)
        values = [
            issue.sequence,
            issue.document,
            issue.page_scope,
            issue.location,
            issue.description,
            issue.resolution,
            issue.proposer,
            issue.status,
        ]
        for column, value in zip([2, 3, 4, 5, 6, 9, 12, 13], values):
            cell = _writable_cell(sheet, row=row, column=column)
            cell.value = value
            cell.alignment = _WRAP_TOP
        sheet.row_dimensions[row].height = 78

    sheet.cell(row=26, column=2).value = '1'
    sheet.cell(row=26, column=3).value = review_plan.pending_problem
    sheet.cell(row=26, column=6).value = review_plan.pending_solution
    sheet.cell(row=26, column=8).value = owner
    sheet.cell(row=26, column=10).value = completion_date
    sheet.cell(row=26, column=12).value = completion_date
    for column in [3, 6, 8, 10, 12]:
        sheet.cell(row=26, column=column).alignment = _WRAP_TOP
    sheet.row_dimensions[26].height = 46

    sheet.cell(row=30, column=5).value = review_plan.conclusion
    sheet.cell(row=31, column=5).value = review_plan.tracking
    sheet.cell(row=32, column=5).value = review_plan.suggestion
    for row in [30, 31, 32]:
        sheet.cell(row=row, column=5).alignment = _WRAP_TOP
        sheet.row_dimensions[row].height = 50


def _clear_issue_rows(sheet) -> None:
    for row in range(13, 23):
        for column in [2, 3, 4, 5, 6, 9, 12, 13]:
            cell = sheet.cell(row=row, column=column)
            if not isinstance(cell, MergedCell):
                cell.value = None
    for row in range(26, 29):
        for column in [2, 3, 6, 8, 10, 12]:
            cell = sheet.cell(row=row, column=column)
            if not isinstance(cell, MergedCell):
                cell.value = None


def _issue_row(sheet, offset: int) -> int:
    # Newer review templates merge each issue record across two rows.
    if isinstance(sheet.cell(row=14, column=9), MergedCell):
        return 13 + (offset * 2)
    return 13 + offset


def _writable_cell(sheet, *, row: int, column: int):
    cell = sheet.cell(row=row, column=column)
    if not isinstance(cell, MergedCell):
        return cell
    for merged_range in sheet.merged_cells.ranges:
        if cell.coordinate in merged_range:
            return sheet.cell(row=merged_range.min_row, column=merged_range.min_col)
    return cell
